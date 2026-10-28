"""
Document Ingestion & Parsing Subsystem.
Supports parsing PDF, DOCX, TXT, and Markdown files from disk or Streamlit uploaded streams.
Tracks page numbers, document names, and extraction metadata.
"""

import io
from pathlib import Path
from typing import List, Union
from dataclasses import dataclass, field

try:
    from langchain_core.documents import Document
except ImportError:
    @dataclass
    class Document:
        page_content: str
        metadata: dict = field(default_factory=dict)


def clean_text(text: str) -> str:
    """Normalize whitespace and strip null bytes."""
    if not text:
        return ""
    # Strip null characters that break certain vector stores
    text = text.replace("\x00", "")
    # Normalize carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return text.strip()


def parse_pdf(file_bytes_or_path: Union[bytes, io.BytesIO, str, Path], filename: str) -> List[Document]:
    """Extract text page-by-page from a PDF document."""
    documents: List[Document] = []
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ImportError("pypdf is required for PDF parsing. Run `pip install pypdf`.")

    stream = file_bytes_or_path if isinstance(file_bytes_or_path, (str, Path)) else (
        file_bytes_or_path if isinstance(file_bytes_or_path, io.BytesIO) else io.BytesIO(file_bytes_or_path)
    )

    reader = PdfReader(stream)
    total_pages = len(reader.pages)

    for page_idx, page in enumerate(reader.pages):
        raw_text = page.extract_text() or ""
        cleaned = clean_text(raw_text)
        if cleaned:
            metadata = {
                "source": filename,
                "page": page_idx + 1,
                "total_pages": total_pages,
                "file_type": "pdf"
            }
            documents.append(Document(page_content=cleaned, metadata=metadata))

    return documents


def parse_docx(file_bytes_or_path: Union[bytes, io.BytesIO, str, Path], filename: str) -> List[Document]:
    """Extract text and tables from a Word (.docx) document."""
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx is required for DOCX parsing. Run `pip install python-docx`.")

    stream = file_bytes_or_path if isinstance(file_bytes_or_path, (str, Path)) else (
        file_bytes_or_path if isinstance(file_bytes_or_path, io.BytesIO) else io.BytesIO(file_bytes_or_path)
    )

    doc = docx.Document(stream)
    paragraphs_text = []

    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            paragraphs_text.append(t)

    # Also capture table contents if present
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                paragraphs_text.append(" | ".join(row_cells))

    combined_text = "\n\n".join(paragraphs_text)
    cleaned = clean_text(combined_text)

    if not cleaned:
        return []

    return [
        Document(
            page_content=cleaned,
            metadata={
                "source": filename,
                "page": 1,
                "total_pages": 1,
                "file_type": "docx"
            }
        )
    ]


def parse_txt(file_bytes_or_path: Union[bytes, io.BytesIO, str, Path], filename: str) -> List[Document]:
    """Extract text from a plain text or markdown file."""
    if isinstance(file_bytes_or_path, (str, Path)):
        with open(file_bytes_or_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
    elif isinstance(file_bytes_or_path, io.BytesIO):
        text = file_bytes_or_path.getvalue().decode("utf-8", errors="replace")
    elif isinstance(file_bytes_or_path, bytes):
        text = file_bytes_or_path.decode("utf-8", errors="replace")
    else:
        text = str(file_bytes_or_path)

    cleaned = clean_text(text)
    if not cleaned:
        return []

    ext = Path(filename).suffix.lower()
    return [
        Document(
            page_content=cleaned,
            metadata={
                "source": filename,
                "page": 1,
                "total_pages": 1,
                "file_type": ext.replace(".", "")
            }
        )
    ]


def load_document(file_source: Union[bytes, io.BytesIO, str, Path], filename: str) -> List[Document]:
    """
    Dispatcher function to parse a document based on file extension.
    Supports PDF, DOCX, TXT, and MD.
    """
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return parse_pdf(file_source, filename)
    elif ext in [".docx", ".doc"]:
        return parse_docx(file_source, filename)
    elif ext in [".txt", ".md", ".log", ".json"]:
        return parse_txt(file_source, filename)
    else:
        # Default fallback to UTF-8 text parser
        return parse_txt(file_source, filename)
